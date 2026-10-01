from __future__ import annotations

import hashlib
import re
from io import BytesIO
from urllib.parse import urljoin

from openpyxl import load_workbook
from bs4 import BeautifulSoup

from app.wind_agents.base import AgentFinding, BaseWindAgent


XLSX_URL = (
    "https://dati.puglia.it/ckan/dataset/"
    "4af29dda-fdcc-4606-bf41-ad4fa3e30790/resource/"
    "1728acca-e2fc-4dcb-bb52-0fc1eaa628c2/download/via_fer.xlsx"
)
SOURCE_URL = "https://dati.puglia.it/ckan/dataset/impianti-proposti-via-fer"
ALBO_URL = "https://albonline.regione.puglia.it/web/guest/home"
ALBO_PREFIX = "_it_linksmt_albopretorio_albopretorio_portlet_AlboPretorioPortlet_"
WIND_TERMS = ("eolico", "eolica", "wind")


class SistemaPugliaWindAgent(BaseWindAgent):
    """Daily Puglia wind discovery from the official VIA FER dataset.

    This replaces the legacy sequential DettaglioInfo id probe. One dataset
    download gives the current regional inventory and makes daily execution
    bounded and deterministic.
    """

    agent_name = "institutional_watch"
    source_name = "Sistema Puglia Energia"
    base_url = SOURCE_URL

    @staticmethod
    def _clean(value: object) -> str:
        return re.sub(r"\s+", " ", str(value or "")).strip()

    @classmethod
    def _header(cls, value: object) -> str:
        text = cls._clean(value).lower()
        for src, dst in (("à", "a"), ("è", "e"), ("é", "e"), ("ì", "i"), ("ò", "o"), ("ù", "u")):
            text = text.replace(src, dst)
        return re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    @classmethod
    def _is_wind(cls, value: object) -> bool:
        text = cls._clean(value).lower()
        return any(term in text for term in WIND_TERMS)

    @classmethod
    def _power_mw(cls, value: object) -> float | None:
        text = cls._clean(value)
        if not text:
            return None
        match = re.search(r"[0-9]+(?:[.,][0-9]+)*", text)
        if not match:
            return None
        raw = match.group(0)
        if "," in raw and "." in raw:
            raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
        elif "," in raw:
            raw = raw.replace(",", ".")
        try:
            power = float(raw)
        except ValueError:
            return None
        return power if 0 < power < 5000 else None

    @classmethod
    def _row_value(cls, row: dict[str, str], *keys: str) -> str | None:
        for key in keys:
            value = cls._clean(row.get(key) or "")
            if value:
                return value
        return None

    @classmethod
    def _external_id(cls, row: dict[str, str], *, province: str, municipality: str, source: str, power_mw: float | None) -> str:
        explicit = cls._row_value(
            row,
            "id", "id_progetto", "id_procedimento", "codice", "codice_pratica",
            "numero_pratica", "numero_procedimento", "procedimento",
        )
        if explicit:
            safe = re.sub(r"[^A-Za-z0-9._-]+", "-", explicit).strip("-")
            return f"PUGLIA-VIA-FER-{safe[:120]}"
        raw = "|".join([province, municipality, source, f"{power_mw:.6f}" if power_mw is not None else ""])
        return "PUGLIA-VIA-FER-" + hashlib.sha1(raw.lower().encode("utf-8")).hexdigest()[:20]

    @classmethod
    def _power_from_text(cls, text: str) -> float | None:
        for match in re.finditer(
            r"(?<![\d.,])([0-9]+(?:[.\s][0-9]{3})*(?:[,\.]\d+)?|[0-9]+(?:[,\.]\d+)?)\s*(MWe|MW)\b",
            text,
            flags=re.I,
        ):
            value = cls._power_mw(match.group(1))
            if value is not None:
                return value
        return None

    @classmethod
    def _proponent_from_text(cls, text: str) -> str | None:
        match = re.search(
            r"Proponente\s*:\s*(.+?)(?=\s+(?:con\s+sede|sede\s+legale|C\.?F\.?|P\.?\s*I(?:VA|va)|codice\s+fiscale)|[.;]|$)",
            text,
            flags=re.I,
        )
        if not match:
            return None
        value = cls._clean(match.group(1)).strip(" -–—:;,.")
        return value if 2 <= len(value) <= 220 else None

    @classmethod
    def _municipalities_from_text(cls, text: str) -> list[str]:
        out: list[str] = []
        patterns = (
            r"(?:nei|ne[Ii]|negli|nel|nella|sito\s+nel|sito\s+nei|da\s+realizzarsi\s+nel|da\s+realizzarsi\s+nei)\s+Comuni?\s+di\s+(.+?)(?:\.|;|,\s*localit[aà]|\s+nonch[eé]|\s+oltre\s+alle|$)",
            r"Comune\s+di\s+([A-ZÀ-Ú][A-Za-zÀ-Úà-ú'’\- ]+?)(?:\s*\([A-Z]{2}\)|,|\.|;|\s+in\s+localit[aà]|$)",
        )
        for pattern in patterns:
            for match in re.finditer(pattern, text, flags=re.I):
                raw = re.sub(r"\([A-Z]{2}\)", "", match.group(1))
                for part in re.split(r",|\s+e\s+|\s+ed\s+", raw, flags=re.I):
                    item = cls._clean(part).strip(" -–—:;,.()")
                    if item and 2 <= len(item) <= 80 and item.lower() not in {x.lower() for x in out}:
                        out.append(item)
                if out:
                    return out[:12]
        return out

    @classmethod
    def _project_name_from_text(cls, text: str, fallback: str) -> str:
        for pattern in (
            r"denominat[oa]\s+[“\"']([^”\"']+)[”\"']",
            r"impianto\s+eolico\s+[“\"']([^”\"']+)[”\"']",
            r"parco\s+eolico\s+[“\"']([^”\"']+)[”\"']",
        ):
            match = re.search(pattern, text, flags=re.I)
            if match:
                return cls._clean(match.group(1))[:900]
        return cls._clean(fallback)[:900]

    def _regional_aoo_id(self) -> str:
        try:
            response = self.session.get(
                ALBO_URL,
                timeout=(8, 20),
                headers={"User-Agent": "Wind-Radar-Agent/0.6"},
            )
            response.raise_for_status()
            soup = BeautifulSoup(response.text, "html.parser")
            select = soup.find("select", attrs={"name": re.compile(r"idAoo", re.I)})
            if select:
                for option in select.find_all("option"):
                    label = self._clean(option.get_text(" ", strip=True)).lower()
                    if "transizione energetica" in label:
                        value = self._clean(option.get("value") or "")
                        if value:
                            return value
        except Exception:
            pass
        return "0"

    @classmethod
    def _regional_external_id(
        cls,
        *,
        project_name: str,
        proponent: str | None,
        municipalities: list[str],
        power_mw: float | None,
        registry: str,
    ) -> str:
        if proponent and municipalities and power_mw:
            raw = "|".join([
                cls._clean(project_name).lower(),
                cls._clean(proponent).lower(),
                "|".join(sorted(cls._clean(x).lower() for x in municipalities)),
                f"{power_mw:.4f}",
            ])
            return "PUGLIA-REGIONAL-WIND-" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:20]
        safe = re.sub(r"[^A-Za-z0-9._-]+", "-", registry).strip("-")
        return f"PUGLIA-REGIONAL-ACT-{safe[:140]}" if safe else "PUGLIA-REGIONAL-ACT-UNKNOWN"

    def _fetch_regional_albo(self, max_pages: int = 12) -> list[AgentFinding]:
        """Regional AU/PAUR/proroga acts from the official Puglia Albo Pretorio."""
        findings: list[AgentFinding] = []
        seen: set[str] = set()
        aoo_id = self._regional_aoo_id()
        for page in range(1, max_pages + 1):
            params = {
                "p_p_id": "it_linksmt_albopretorio_albopretorio_portlet_AlboPretorioPortlet",
                "p_p_lifecycle": "0",
                "p_p_state": "normal",
                "p_p_mode": "view",
                f"{ALBO_PREFIX}mvcRenderCommandName": "/cercaAtto",
                f"{ALBO_PREFIX}cur": str(page),
                f"{ALBO_PREFIX}delta": "60",
                f"{ALBO_PREFIX}idAoo": aoo_id,
                f"{ALBO_PREFIX}idStatoAtto": "-1",
                f"{ALBO_PREFIX}idTipoAtto": "0",
                f"{ALBO_PREFIX}resetCur": "false",
            }
            response = self.session.get(
                ALBO_URL,
                params=params,
                timeout=(8, 25),
                headers={"User-Agent": "Wind-Radar-Agent/0.6"},
            )
            response.raise_for_status()
            soup = BeautifulSoup(response.text, "html.parser")
            rows = soup.find_all("tr")
            if not rows:
                break
            page_hits = 0
            for tr in rows:
                cells = tr.find_all("td")
                if len(cells) < 6:
                    continue
                values = [self._clean(td.get_text(" ", strip=True)) for td in cells]
                registry = values[0] if values else ""
                adoption_number = values[2] if len(values) > 2 else ""
                adoption_date = values[3] if len(values) > 3 else ""
                structure = values[4] if len(values) > 4 else ""
                object_text = values[5] if len(values) > 5 else ""
                combined = self._clean(f"{structure} {object_text}")
                if not self._is_wind(combined):
                    continue
                lowered = combined.lower()
                if not any(token in lowered for token in ("autorizzazione unica", "p.a.u.r", "paur", "proroga")):
                    continue

                proponent = self._proponent_from_text(object_text)
                power_mw = self._power_from_text(object_text)
                municipalities = self._municipalities_from_text(object_text)
                source_anchor = tr.find("a", href=True)
                source_url = urljoin(ALBO_URL, source_anchor.get("href")) if source_anchor else str(response.url)
                project_name = self._project_name_from_text(object_text, object_text)
                stable = registry or f"{adoption_number}-{adoption_date}-{hashlib.sha1(object_text.encode('utf-8')).hexdigest()[:10]}"
                external_id = self._regional_external_id(
                    project_name=project_name,
                    proponent=proponent,
                    municipalities=municipalities,
                    power_mw=power_mw,
                    registry=stable,
                )
                if external_id in seen:
                    continue
                seen.add(external_id)
                page_hits += 1

                procedure = "Proroga AU/PAUR" if "proroga" in lowered else "Autorizzazione Unica / PAUR"
                findings.append(
                    AgentFinding(
                        external_id=external_id,
                        source_name="Regione Puglia AU/PAUR",
                        source_url=source_url,
                        title=project_name,
                        finding_type="project_source",
                        payload={
                            "project_name": project_name,
                            "proponent": proponent,
                            "region": "Puglia",
                            "province": None,
                            "municipalities": municipalities,
                            "power_mw": power_mw,
                            "procedure": procedure,
                            "status_raw": object_text[:1200],
                            "source_date": adoption_date or None,
                            "albo_registry": registry or None,
                            "adoption_number": adoption_number or None,
                            "sector": "eolico",
                            "source_grade_ceiling": "A1",
                            "project_specific": True,
                            "source_adapter_origin": "regional_puglia_albo",
                            "ingestion_path": "official_regional_albo_au_paur",
                        },
                    )
                )
            # Keep paging within the bounded window: the Albo may interleave
            # unrelated acts even when the AOO filter is unavailable.
        return findings

    def _fetch_mase_via_dataset(self) -> list[AgentFinding]:
        response = self.session.get(
            XLSX_URL,
            timeout=(8, 30),
            headers={"User-Agent": "Wind-Radar-Agent/0.6"},
        )
        response.raise_for_status()
        wb = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if not raw_headers:
            return []
        headers = [self._header(value) for value in raw_headers]

        findings: list[AgentFinding] = []
        seen: set[str] = set()
        for values in iterator:
            row = {
                header: self._clean(values[index]) if index < len(values) and values[index] is not None else ""
                for index, header in enumerate(headers)
                if header
            }
            source = self._row_value(row, "fonte", "tipologia_fonte", "tecnologia") or ""
            if not self._is_wind(source):
                continue

            province = self._row_value(row, "provincia") or ""
            municipality = self._row_value(row, "comune", "municipio") or ""
            power_mw = self._power_mw(self._row_value(row, "potenza_mw", "potenza", "potenza_nominale"))
            status = self._row_value(row, "stato_del_procedimento", "stato", "esito")
            proponent = self._row_value(row, "proponente", "societa_proponente", "soggetto_proponente")
            project_name = self._row_value(row, "denominazione", "progetto", "oggetto", "nome_progetto")
            if not project_name:
                area = municipality or "Comune non indicato"
                if province:
                    area = f"{area} ({province})"
                project_name = f"Impianto eolico - {area}" + (f" - {power_mw:g} MW" if power_mw else "")

            external_id = self._external_id(
                row,
                province=province,
                municipality=municipality,
                source=source,
                power_mw=power_mw,
            )
            if external_id in seen:
                continue
            seen.add(external_id)

            findings.append(
                AgentFinding(
                    external_id=external_id,
                    source_name=self.source_name,
                    source_url=XLSX_URL,
                    title=project_name[:900],
                    finding_type="project_source",
                    payload={
                        "project_name": project_name[:900],
                        "proponent": proponent,
                        "region": "Puglia",
                        "province": province or None,
                        "municipalities": [municipality] if municipality else [],
                        "power_mw": power_mw,
                        "procedure": "VIA FER",
                        "status_raw": status,
                        "source_type": source,
                        "sector": "eolico",
                        "source_grade_ceiling": "A1",
                        "project_specific": True,
                        "source_adapter_origin": "pv_agent_mvp/puglia.py",
                        "ingestion_path": "official_via_fer_xlsx",
                    },
                )
            )

        return findings

    def fetch(self) -> list[AgentFinding]:
        return self._fetch_mase_via_dataset()


class PugliaRegionalAuWindAgent(SistemaPugliaWindAgent):
    """Independent regional AU/PAUR source with its own bootstrap lifecycle."""

    source_name = "Regione Puglia AU/PAUR"
    base_url = ALBO_URL

    def fetch(self) -> list[AgentFinding]:
        return self._fetch_regional_albo()
